"""
Document extraction: ZIP -> list[ParsedDoc]. No LLM calls here.

Entry point: extract_zip(zip_path) -> list[ParsedDoc]

Supported formats: PDF, DOCX, XLSX, CSV, TXT.
Unsupported files are skipped with a warning (no crash).
"""

import io
import os
import re
import tempfile
import zipfile

import fitz                         # PyMuPDF — PDF parsing
import pandas as pd
import pytesseract
from PIL import Image
from docx import Document           # python-docx — DOCX parsing

# Point pytesseract at the Tesseract executable if it's at the Windows default
# install path; otherwise leave pytesseract to resolve it from PATH (so the app
# still OCRs on machines with a different install location).
_TESSERACT_WIN_DEFAULT = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
if os.path.exists(_TESSERACT_WIN_DEFAULT):
    pytesseract.pytesseract.tesseract_cmd = _TESSERACT_WIN_DEFAULT

from core.schema import ParsedDoc

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".txt"}

OCR_DPI = 300   # higher DPI = sharper image = more accurate OCR (200 was missing fine print)


# ── OCR helpers ───────────────────────────────────────────────────────────────

def _page_to_image(page) -> "Image.Image":
    """
    Render a PDF page to a PIL image robustly. Going via PNG bytes means we
    don't assume a colour space — works for RGB, RGBA, CMYK and greyscale pages
    (the old RGB-only frombytes path crashed on some PDFs).
    """
    pix = page.get_pixmap(dpi=OCR_DPI)
    return Image.open(io.BytesIO(pix.tobytes("png")))


def _norm_line(s: str) -> str:
    """Lowercase + keep only alphanumerics — used to dedupe OCR lines across passes."""
    return re.sub(r"[^a-z0-9]", "", s.lower())


# Supplemental OCR only runs when images cover at least this fraction of the
# page. Header logos cover ~1-2% and contain no content — OCRing them wasted a
# dual Tesseract pass on almost every page of a typical report (measured: 21 of
# 22 pages on the sample report had a logo; only 4 had a real embedded scan).
_OCR_MIN_IMAGE_FRACTION = 0.05


def _image_area_fraction(page) -> float:
    """Fraction (0..1) of the page area covered by embedded images."""
    page_area = page.rect.width * page.rect.height
    if page_area <= 0:
        return 0.0
    covered = 0.0
    for img in page.get_images(full=True):
        try:
            for r in page.get_image_rects(img[0]):
                covered += r.width * r.height
        except Exception:
            continue
    return min(covered / page_area, 1.0)


def _ocr_image(img) -> str:
    """
    Run OCR with two Tesseract page-segmentation modes and merge them:
      - psm 3  (default): good reading order for flowing paragraphs
      - psm 11 (sparse) : catches stray text the layout pass misses
                          (e.g. a date in a top corner next to a logo)
    We keep psm 3 as the base for readability, then append any lines psm 11
    found that psm 3 didn't. Grayscale first — it measurably helps Tesseract.
    """
    gray = img.convert("L")
    primary = pytesseract.image_to_string(gray, config="--psm 3").strip()
    sparse = pytesseract.image_to_string(gray, config="--psm 11").strip()

    seen = {_norm_line(ln) for ln in primary.splitlines() if ln.strip()}
    extras = [
        ln.strip()
        for ln in sparse.splitlines()
        if ln.strip() and _norm_line(ln) and _norm_line(ln) not in seen
    ]
    if extras:
        return (primary + "\n" + "\n".join(extras)).strip()
    return primary


# ── ZIP handler ───────────────────────────────────────────────────────────────

def extract_zip(zip_path: str) -> list[ParsedDoc]:
    """
    Open a ZIP file, parse every supported file inside it, and return
    a list of ParsedDoc objects — one per supported file found.
    """
    results = []

    with zipfile.ZipFile(zip_path, "r") as zf:
        entries = [
            name for name in zf.namelist()  
            if not name.endswith("/")
            and not os.path.basename(name).startswith(".")
        ]

        seen_names = set()
        for entry in entries:
            ext = os.path.splitext(entry)[1].lower()
            if ext not in SUPPORTED_EXTENSIONS:
                print(f"  [skip] {entry} — unsupported format")
                continue

            file_bytes = zf.read(entry)
            filename = os.path.basename(entry)

            # Two files with the same name in different ZIP folders would collapse
            # into one label (breaking per-document attribution and filtering) —
            # disambiguate with the parent folder name.
            if filename in seen_names:
                parent = os.path.basename(os.path.dirname(entry))
                filename = f"{parent}/{filename}" if parent else f"copy_{filename}"
            seen_names.add(filename)

            print(f"  [parse] {filename} ({ext})")
            try:
                doc = _parse_file(filename, ext, file_bytes)
                results.append(doc)
            except Exception as e:
                print(f"  [error] {filename}: {e}")
                results.append(ParsedDoc(
                    filename=filename,
                    doc_type=ext.lstrip("."),
                    raw_text="",
                    tables=[],
                    warnings=[f"Parsing failed: {e}"],
                ))

    return results


def list_zip_contents(zip_path: str) -> list[dict]:
    """
    Return a summary of files inside a ZIP without parsing them.
    Used by Page 1 to show the user what's in their upload before processing.
    """
    contents = []
    with zipfile.ZipFile(zip_path, "r") as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            ext = os.path.splitext(info.filename)[1].lower()
            contents.append({
                "name": os.path.basename(info.filename),
                "extension": ext,
                "size_kb": round(info.file_size / 1024, 1),
                "supported": ext in SUPPORTED_EXTENSIONS,
            })
    return contents


# ── format dispatcher ─────────────────────────────────────────────────────────

def _parse_file(filename: str, ext: str, file_bytes: bytes) -> ParsedDoc:
    if ext == ".pdf":
        return _parse_pdf(filename, file_bytes)
    elif ext == ".docx":
        return _parse_docx(filename, file_bytes)
    elif ext == ".xlsx":
        return _parse_xlsx(filename, file_bytes)
    elif ext == ".csv":
        return _parse_csv(filename, file_bytes)
    elif ext == ".txt":
        return _parse_txt(filename, file_bytes)
    else:
        raise ValueError(f"No parser for extension: {ext}")


# ── per-format parsers ────────────────────────────────────────────────────────

def _parse_pdf(filename: str, file_bytes: bytes) -> ParsedDoc:
    warnings = []
    text_parts = []
    tables = []

    pdf = fitz.open(stream=file_bytes, filetype="pdf")
    page_count = len(pdf)

    for page_num, page in enumerate(pdf, start=1):
        page_text = page.get_text("text").strip()

        if not page_text:
            # Page is a scanned image — run OCR via Tesseract
            try:
                page_text = _ocr_image(_page_to_image(page))
                if page_text:
                    warnings.append(f"Page {page_num} was a scanned image — OCR applied.")
                else:
                    warnings.append(f"Page {page_num}: OCR returned no text.")
                    continue
            except Exception as e:
                warnings.append(f"Page {page_num}: OCR failed — {e}")
                continue
        elif _image_area_fraction(page) >= _OCR_MIN_IMAGE_FRACTION:
            # Page has clean digital text but also a SUBSTANTIAL embedded image
            # (e.g. a scanned signed letter alongside typed text — not just a
            # header logo). OCR the page and APPEND the result — we keep the
            # accurate digital text and add whatever lived only inside the image.
            try:
                ocr_text = _ocr_image(_page_to_image(page))
                if len(ocr_text.split()) > len(page_text.split()) * 1.5:
                    page_text = page_text + "\n[Image content via OCR]\n" + ocr_text
                    warnings.append(f"Page {page_num}: embedded image detected — OCR appended to capture full content.")
            except Exception as e:
                warnings.append(f"Page {page_num}: supplemental OCR failed — {e}")

        lines = [ln for ln in page_text.splitlines() if ln.strip()]
        text_parts.append(f"[Page {page_num}]\n" + "\n".join(lines))

        for table in page.find_tables():
            df = table.to_pandas()
            if not df.empty:
                # Replace empty cells (pandas NaN) with "" so they don't render as
                # the literal noise token "nan" when we stringify the table.
                df = df.fillna("").astype(str)

                # PyMuPDF names columns "Col0", "Col1", … (or "0", "1", …) when it
                # can't detect a header row — the table's REAL header is then
                # sitting in the first data row. Promote it so the table keeps its
                # actual column names instead of numeric placeholders.
                cols = [str(c) for c in df.columns]
                if len(df) > 1 and all(re.fullmatch(r"(Col)?\d+", c) for c in cols):
                    df.columns = [str(v).strip() or f"Column {i + 1}"
                                  for i, v in enumerate(df.iloc[0])]
                    df = df.iloc[1:].reset_index(drop=True)

                tables.append(_df_to_dict(df, source=f"{filename} p{page_num}"))

                # Render the table as clearly labelled text so Gemini understands
                # its structure — columns, and each row with its values explicitly stated.
                #[TABLE from report.pdf page 4]
                #Columns: Finding | Severity | Status
                #Row 1: Finding: Missing logs | Severity: High | Status: Open
                #Row 2: Finding: Expense gap | Severity: Medium | Status: Closed
                #[END TABLE]
                col_names = " | ".join(df.columns.tolist())
                table_lines = [
                    f"[TABLE from {filename} page {page_num}]",
                    f"Columns: {col_names}",
                ]
                for row_idx, row in df.iterrows():
                    row_text = " | ".join(f"{col}: {val}" for col, val in row.items())
                    table_lines.append(f"Row {row_idx + 1}: {row_text}")
                table_lines.append("[END TABLE]")
                text_parts.append("\n".join(table_lines))

    pdf.close()

    return ParsedDoc(
        filename=filename,
        doc_type="pdf",
        raw_text="\n\n".join(text_parts),
        tables=tables,
        page_count=page_count,
        warnings=warnings,
    )


def _parse_docx(filename: str, file_bytes: bytes) -> ParsedDoc:
    doc = Document(io.BytesIO(file_bytes))
    text_parts = []
    tables = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            text_parts.append(text)

    for i, table in enumerate(doc.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            headers = rows[0]
            data_rows = rows[1:]
            tables.append({
                "source": f"{filename} table {i}",
                "headers": headers,
                "rows": data_rows,
            })
            table_lines = [" | ".join(headers)]
            table_lines += [" | ".join(r) for r in data_rows]
            text_parts.append("\n".join(table_lines))

    return ParsedDoc(
        filename=filename,
        doc_type="docx",
        raw_text="\n\n".join(text_parts),
        tables=tables,
    )


def _parse_xlsx(filename: str, file_bytes: bytes) -> ParsedDoc:
    text_parts = []
    tables = []

    xl = pd.ExcelFile(io.BytesIO(file_bytes))
    for sheet_name in xl.sheet_names:
        df = xl.parse(sheet_name).fillna("")

        if df.empty:
            continue

        tables.append(_df_to_dict(df, source=f"{filename} [{sheet_name}]"))
        text_parts.append(f"=== Sheet: {sheet_name} ===")
        text_parts.append(df.to_string(index=False))

    return ParsedDoc(
        filename=filename,
        doc_type="xlsx",
        raw_text="\n\n".join(text_parts),
        tables=tables,
    )


def _parse_csv(filename: str, file_bytes: bytes) -> ParsedDoc:
    try:
        df = pd.read_csv(io.BytesIO(file_bytes)).fillna("")
    except UnicodeDecodeError:
        # Excel commonly exports CSVs as cp1252/latin-1 rather than UTF-8.
        df = pd.read_csv(io.BytesIO(file_bytes), encoding="latin-1").fillna("")
    return ParsedDoc(
        filename=filename,
        doc_type="csv",
        raw_text=df.to_string(index=False),
        tables=[_df_to_dict(df, source=filename)],
    )


def _parse_txt(filename: str, file_bytes: bytes) -> ParsedDoc:
    raw_text = file_bytes.decode("utf-8", errors="replace").strip()
    return ParsedDoc(
        filename=filename,
        doc_type="txt",
        raw_text=raw_text,
        tables=[],
    )


# ── helpers ───────────────────────────────────────────────────────────────────

def _df_to_dict(df: pd.DataFrame, source: str) -> dict:
    df = df.astype(str)
    return {
        "source": source,
        "headers": df.columns.tolist(),
        "rows": df.values.tolist(),
    }
